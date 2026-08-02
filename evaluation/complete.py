"""Avaliação completa, auditável e reproduzível do benchmark RAG 2.0."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import statistics
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import chromadb
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from app.config import get_settings
from app.rag.factory import create_llm_provider
from evaluation.io import read_json


JUDGE_VERSION = "1.0"
JUDGE_SYSTEM_PROMPT = """Você é um avaliador acadêmico conservador de respostas RAG.
Avalie somente com base no gabarito fixo e nas evidências fornecidas. Nunca altere,
corrija ou amplie o gabarito. Não considere uma resposta correta por mera semelhança
textual. Exija sustentação para cada afirmação verificável. Se houver ambiguidade,
marque revisao_manual_necessaria=true. Retorne somente JSON válido, sem markdown.
"""
JUDGE_USER_TEMPLATE = """Avalie esta resposta.

PERGUNTA:
{question}

RESPOSTA ESPERADA FIXA:
{expected_answer}

FATOS OBRIGATÓRIOS (preserve a ordem):
{required_facts}

EVIDÊNCIA VALIDADA DO GABARITO:
{expected_evidence}

RESPOSTA GERADA:
{produced_answer}

FONTES CITÁVEIS DO CONTEXTO (o número corresponde a [Fonte N]):
{cited_sources}

RESULTADOS BRUTOS DA RECUPERAÇÃO:
{retrieval_results}

Retorne exatamente um objeto JSON com:
- factual_correct: 0 ou 1;
- facts_mentioned: lista de booleanos, um por fato obrigatório e na mesma ordem;
- faithfulness: 0, 0.5 ou 1;
- relevance: 0, 0.5 ou 1;
- citation_support: 0, 0.5 ou 1;
- supporting_citation_indices: lista de inteiros correspondentes a [Fonte N];
- supporting_retrieval_indices: lista de inteiros correspondentes a [Resultado N];
- justification: justificativa curta e específica;
- evidence_used: trecho ou descrição da evidência usada;
- ambiguous: booleano;
- manual_review_reason: texto, vazio somente se não houver dúvida.
"""

CITATION_RE = re.compile(r"\[Fonte\s+(\d+)\]", re.IGNORECASE)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _normalized_name(value: str | None) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).casefold().strip()
    return re.sub(r"\s+", " ", text)


def _same_document(left: str | None, right: str | None) -> bool:
    return bool(left and right and _normalized_name(left) == _normalized_name(right))


def _pages(result: dict[str, Any]) -> set[int]:
    start = result.get("page_start")
    end = result.get("page_end") or start
    if not isinstance(start, int) or not isinstance(end, int) or end < start:
        return set()
    return set(range(start, end + 1))


def _page_hit(
    results: list[dict[str, Any]],
    expected_document: str | None,
    expected_pages: list[int],
    tolerance: int = 0,
) -> bool:
    acceptable = {
        page + offset
        for page in expected_pages
        for offset in range(-tolerance, tolerance + 1)
        if page + offset >= 1
    }
    return any(
        _same_document(item.get("file_name"), expected_document)
        and bool(_pages(item) & acceptable)
        for item in results
    )


def _percentile(values: list[float], fraction: float) -> float | None:
    if not values:
        return None
    ordered = sorted(values)
    return float(ordered[max(0, math.ceil(fraction * len(ordered)) - 1)])


def _latency(values: list[float]) -> dict[str, float | None]:
    return {
        "mean": statistics.fmean(values) if values else None,
        "median": statistics.median(values) if values else None,
        "minimum": min(values) if values else None,
        "maximum": max(values) if values else None,
        "standard_deviation": statistics.pstdev(values) if values else None,
        "p95": _percentile(values, 0.95),
        "p99": _percentile(values, 0.99),
    }


def _extract_json(text: str) -> dict[str, Any]:
    value = text.strip()
    if value.startswith("```"):
        value = re.sub(r"^```(?:json)?\s*|\s*```$", "", value, flags=re.I)
    start, end = value.find("{"), value.rfind("}")
    if start < 0 or end < start:
        raise ValueError("juiz não retornou objeto JSON")
    parsed = json.loads(value[start : end + 1])
    if not isinstance(parsed, dict):
        raise ValueError("saída do juiz não é objeto")
    return parsed


def _source_text(source: dict[str, Any], index: int) -> str:
    content = str(source.get("content") or "")[:1800]
    return (
        f"[Fonte {index}] arquivo={source.get('file_name')}; "
        f"páginas={source.get('page_start')}-{source.get('page_end')}; "
        f"seção={source.get('section')}; trecho={content}"
    )


def _result_text(result: dict[str, Any], index: int) -> str:
    content = str(result.get("content") or "")[:1000]
    return (
        f"[Resultado {index}] arquivo={result.get('file_name')}; "
        f"páginas={result.get('page_start')}-{result.get('page_end')}; "
        f"score={result.get('score')}; trecho={content}"
    )


def _hydrate_sources(records: list[dict[str, Any]]) -> None:
    settings = get_settings()
    client = chromadb.PersistentClient(path=str(settings.chroma_dir))
    collection = client.get_collection(settings.chroma_collection)
    ids = sorted(
        {
            str(source.get("chunk_id"))
            for record in records
            for source in record.get("sources", [])
            if source.get("chunk_id")
        }
    )
    contents: dict[str, str] = {}
    for offset in range(0, len(ids), 500):
        response = collection.get(ids=ids[offset : offset + 500], include=["documents"])
        for chunk_id, content in zip(
            response.get("ids", []), response.get("documents") or [], strict=False
        ):
            contents[str(chunk_id)] = str(content or "")
    for record in records:
        record["sources"] = [
            {**source, "content": contents.get(str(source.get("chunk_id")), "")}
            for source in record.get("sources", [])
        ]


def _judge_one(
    provider: Any,
    benchmark: dict[str, Any],
    record: dict[str, Any],
) -> tuple[dict[str, Any], str]:
    prompt = JUDGE_USER_TEMPLATE.format(
        question=benchmark["question"],
        expected_answer=benchmark["expected_answer"],
        required_facts=json.dumps(benchmark.get("required_facts", []), ensure_ascii=False),
        expected_evidence=benchmark.get("evidence", ""),
        produced_answer=record.get("produced_answer", ""),
        cited_sources="\n".join(
            _source_text(source, index)
            for index, source in enumerate(record.get("sources", []), 1)
        )
        or "Nenhuma fonte.",
        retrieval_results="\n".join(
            _result_text(result, index)
            for index, result in enumerate(record.get("retrieval_results", []), 1)
        )
        or "Nenhum resultado.",
    )
    raw = provider.generate(
        system_prompt=JUDGE_SYSTEM_PROMPT,
        user_prompt=prompt,
        timeout_seconds=90,
        temperature=0,
        top_p=1,
        max_tokens=1200,
    )
    return _extract_json(raw), raw


def _valid_score(value: Any, allowed: set[float]) -> float | None:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if number in allowed else None


def _citation_indices(answer: str, source_count: int) -> list[int]:
    return list(
        dict.fromkeys(
            index
            for index in (int(value) for value in CITATION_RE.findall(answer))
            if 1 <= index <= source_count
        )
    )


def _derived_record(
    benchmark: dict[str, Any],
    record: dict[str, Any],
    judge: dict[str, Any] | None,
    raw_judge: str | None,
    judge_error: str | None,
    repetition: int,
) -> dict[str, Any]:
    answerable = bool(benchmark["answerable"])
    stage_one_results = list(record.get("retrieval_results", []))
    context_results = list(
        record.get("context_results") or record.get("sources", [])
    )
    results = context_results[: int(record.get("top_k", 0))]
    expected_document = benchmark.get("expected_document")
    expected_pages = list(benchmark.get("expected_pages", []))
    documents = [item.get("file_name") for item in results]
    rank = next(
        (index for index, name in enumerate(documents, 1) if _same_document(name, expected_document)),
        None,
    )
    citations = _citation_indices(
        str(record.get("produced_answer") or ""), len(record.get("sources", []))
    )
    cited_sources = [record["sources"][index - 1] for index in citations]
    source_present = bool(citations)
    page_present = any(isinstance(source.get("page_start"), int) for source in cited_sources)
    factual = _valid_score((judge or {}).get("factual_correct"), {0, 1}) if answerable else None
    faithfulness = _valid_score((judge or {}).get("faithfulness"), {0, 0.5, 1}) if answerable else None
    relevance = _valid_score((judge or {}).get("relevance"), {0, 0.5, 1}) if answerable else None
    citation_support = _valid_score((judge or {}).get("citation_support"), {0, 0.5, 1}) if answerable else None
    facts = (judge or {}).get("facts_mentioned", [])
    facts_valid = (
        isinstance(facts, list)
        and len(facts) == len(benchmark.get("required_facts", []))
        and all(isinstance(value, bool) for value in facts)
    )
    completeness = (
        sum(facts) / len(facts)
        if answerable and facts_valid and facts
        else (1.0 if answerable and facts_valid and not facts else None)
    )
    supporting_citations = {
        int(value)
        for value in (judge or {}).get("supporting_citation_indices", [])
        if isinstance(value, int) and 1 <= value <= len(record.get("sources", []))
    }
    supporting_results = {
        int(value)
        for value in (judge or {}).get("supporting_retrieval_indices", [])
        if isinstance(value, int) and 1 <= value <= len(results)
    }
    supported_cited_sources = [
        record["sources"][index - 1]
        for index in supporting_citations & set(citations)
    ]
    citation_document_correct = (
        None
        if answerable and citation_support is None
        else bool(answerable and citation_support == 1 and supported_cited_sources)
    )
    citation_page_exact = (
        None
        if answerable and citation_support is None
        else bool(
            citation_document_correct
            and any(
                _same_document(source.get("file_name"), expected_document)
                and bool(_pages(source) & set(expected_pages))
                for source in supported_cited_sources
            )
        )
    )
    citation_page_evidence = (
        None
        if answerable and citation_support is None
        else bool(citation_document_correct)
    )
    retrieval_evidence_valid = (
        None
        if answerable and judge is None
        else bool(answerable and supporting_results)
    )
    quality_scores = (
        factual,
        completeness,
        faithfulness,
        relevance,
        citation_support,
    )
    fully_correct = (
        None
        if answerable and any(score is None for score in quality_scores)
        else bool(
            answerable
            and not record.get("refused")
            and factual == 1
            and completeness == 1
            and faithfulness == 1
            and relevance == 1
            and citation_support == 1
        )
    )
    ambiguous = bool((judge or {}).get("ambiguous"))
    deterministic_audit_sample = answerable and int(re.sub(r"\D", "", benchmark["id"]) or 0) % 5 == 0
    manual = bool(
        answerable
        and (
            judge_error
            or ambiguous
            or not fully_correct
            or deterministic_audit_sample
        )
    )
    return {
        "id": benchmark["id"],
        "repetition": repetition,
        "split": benchmark["split"],
        "question": benchmark["question"],
        "answerable": answerable,
        "question_type": benchmark["question_type"],
        "difficulty": benchmark["difficulty"],
        "expected_answer": benchmark["expected_answer"],
        "required_facts": benchmark.get("required_facts", []),
        "expected_evidence": benchmark.get("evidence", ""),
        "produced_answer": record.get("produced_answer", ""),
        "expected_document": expected_document,
        "expected_pages": expected_pages,
        "printed_pages": benchmark.get("printed_pages", []),
        "retrieval_results": results,
        "stage_one_retrieval_results": stage_one_results,
        "context_results": context_results,
        "retrieved_documents": documents,
        "retrieved_pages": record.get("retrieved_pages", []),
        "scores": [item.get("score") for item in results],
        "positions": list(range(1, len(results) + 1)),
        "sources": record.get("sources", []),
        "citation_indices": citations,
        "hit_rate_at_1": bool(answerable and rank == 1),
        "hit_rate_at_k": bool(answerable and rank is not None),
        "first_relevant_rank": rank,
        "reciprocal_rank": 1 / rank if answerable and rank else 0.0,
        "page_exact": bool(answerable and _page_hit(results, expected_document, expected_pages)),
        "page_tolerance_1": bool(answerable and _page_hit(results, expected_document, expected_pages, 1)),
        "page_evidence_valid": retrieval_evidence_valid if answerable else None,
        "citation_document_present": source_present,
        "citation_document_correct": citation_document_correct if answerable else None,
        "citation_page_present": page_present,
        "citation_page_exact": citation_page_exact if answerable else None,
        "citation_page_evidence_valid": citation_page_evidence if answerable else None,
        "factual_correct": factual,
        "completeness": completeness,
        "faithfulness": faithfulness,
        "relevance": relevance,
        "citation_support": citation_support,
        "fully_correct": fully_correct if answerable else None,
        "refused": bool(record.get("refused")),
        "correct_refusal": bool(not answerable and record.get("refused")),
        "improper_refusal": bool(answerable and record.get("refused")),
        "improper_answer": bool(not answerable and not record.get("refused")),
        "total_time_ms": record.get("total_time_ms"),
        "retrieval_time_ms": record.get("retrieval_time_ms"),
        "generation_time_ms": record.get("generation_time_ms"),
        "empty_answer": not bool(str(record.get("produced_answer") or "").strip()),
        "error": record.get("error"),
        "timeout": bool(record.get("timeout")),
        "judge_model": get_settings().llm_model,
        "judge_version": JUDGE_VERSION,
        "judge_temperature": 0,
        "judge_raw": raw_judge,
        "judge_error": judge_error,
        "judge_justification": (judge or {}).get("justification", ""),
        "judge_evidence_used": (judge or {}).get("evidence_used", ""),
        "revisao_manual_necessaria": manual,
        "manual_review_reason": (
            judge_error
            or (judge or {}).get("manual_review_reason")
            or ("Amostra sistemática de 20% para validação humana." if deterministic_audit_sample else "")
        ),
        "settings": record.get("settings", {}),
        "raw_record": record,
    }


def _metric(
    rows: list[dict[str, Any]],
    name: str,
    values: list[bool | float | int],
    *,
    kind: str = "rate",
    notes: str = "",
) -> dict[str, Any]:
    measurable = [value for value in values if value is not None]
    if not measurable:
        value = numerator = None
        denominator = 0
    elif kind == "rate":
        numerator = sum(bool(value) for value in measurable)
        denominator = len(measurable)
        value = numerator / denominator
    else:
        numerator = None
        denominator = len(measurable)
        value = statistics.fmean(float(item) for item in measurable)
    return {
        "metric": name,
        "value": value,
        "numerator": numerator,
        "denominator": denominator,
        "percentage": value * 100 if value is not None and kind == "rate" else None,
        "measurable": bool(measurable),
        "notes": notes,
    }


def _scalar_metric(
    name: str,
    value: float | int,
    denominator: int,
    *,
    percentage: bool = False,
    notes: str = "",
) -> dict[str, Any]:
    return {
        "metric": name,
        "value": value,
        "numerator": None,
        "denominator": denominator,
        "percentage": value * 100 if percentage else None,
        "measurable": True,
        "notes": notes,
    }


def calculate_summary(rows: list[dict[str, Any]], scope: str) -> list[dict[str, Any]]:
    selected = rows if scope == "overall" else [row for row in rows if row["split"] == scope]
    answerable = [row for row in selected if row["answerable"]]
    negative = [row for row in selected if not row["answerable"]]
    answered = [row for row in selected if not row["refused"]]
    true_positives = sum(row["answerable"] and not row["refused"] for row in selected)
    false_positives = sum(not row["answerable"] and not row["refused"] for row in selected)
    false_negatives = sum(row["answerable"] and row["refused"] for row in selected)
    true_negatives = sum(not row["answerable"] and row["refused"] for row in selected)
    precision = true_positives / len(answered) if answered else 0.0
    recall = true_positives / len(answerable) if answerable else 0.0
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
    accuracy = (true_positives + true_negatives) / len(selected) if selected else 0.0
    always_answer_accuracy = len(answerable) / len(selected) if selected else 0.0
    k = int(selected[0]["settings"].get("top_k", 0)) if selected else 0
    metrics = [
        _metric(answerable, "hit_rate_at_1", [row["hit_rate_at_1"] for row in answerable]),
        _metric(answerable, f"hit_rate_at_{k}", [row["hit_rate_at_k"] for row in answerable]),
        _metric(answerable, "mrr", [row["reciprocal_rank"] for row in answerable], kind="mean"),
        _metric(answerable, "page_exact", [row["page_exact"] for row in answerable]),
        _metric(answerable, "page_tolerance_1", [row["page_tolerance_1"] for row in answerable]),
        _metric(answerable, "page_evidence_valid_llm", [row["page_evidence_valid"] for row in answerable], notes="Critério assistido por juiz LLM; revisar casos marcados."),
        _metric(answerable, "citation_document_present", [row["citation_document_present"] for row in answerable]),
        _metric(answerable, "citation_document_correct", [row["citation_document_correct"] for row in answerable]),
        _metric(answerable, "citation_page_present", [row["citation_page_present"] for row in answerable]),
        _metric(answerable, "citation_page_exact", [row["citation_page_exact"] for row in answerable]),
        _metric(answerable, "citation_page_evidence_valid_llm", [row["citation_page_evidence_valid"] for row in answerable], notes="Critério assistido por juiz LLM; revisar casos marcados."),
        _metric(answerable, "factual_correct", [row["factual_correct"] for row in answerable]),
        _metric(answerable, "mean_completeness", [row["completeness"] for row in answerable], kind="mean"),
        _metric(answerable, "mean_faithfulness", [row["faithfulness"] for row in answerable], kind="mean"),
        _metric(answerable, "mean_relevance", [row["relevance"] for row in answerable], kind="mean"),
        _metric(answerable, "fully_correct", [row["fully_correct"] for row in answerable]),
        _metric(negative, "correct_refusal", [row["correct_refusal"] for row in negative]),
        _metric(answerable, "improper_refusal", [row["improper_refusal"] for row in answerable]),
        _metric(negative, "improper_answer", [row["improper_answer"] for row in negative]),
        _scalar_metric("true_positives", true_positives, len(selected)),
        _scalar_metric("false_positives", false_positives, len(selected)),
        _scalar_metric("false_negatives", false_negatives, len(selected)),
        _scalar_metric("true_negatives", true_negatives, len(selected)),
        _scalar_metric("decision_accuracy", accuracy, len(selected), percentage=True),
        _scalar_metric("decision_precision", precision, len(answered), percentage=True),
        _scalar_metric("decision_recall", recall, len(answerable), percentage=True),
        _scalar_metric("decision_f1", f1, len(selected), percentage=True),
        _scalar_metric(
            "always_answer_accuracy",
            always_answer_accuracy,
            len(selected),
            percentage=True,
            notes="Baseline trivial que responde a todas as perguntas.",
        ),
        _metric(
            selected,
            "valid_inline_source",
            [bool(row["citation_indices"]) for row in selected],
        ),
        _metric(
            negative,
            "correct_refusal_without_source",
            [
                bool(
                    row["correct_refusal"]
                    and not row["sources"]
                    and not row["citation_indices"]
                )
                for row in negative
            ],
        ),
        _metric(selected, "empty_answer", [row["empty_answer"] for row in selected]),
        _metric(selected, "errors", [bool(row["error"]) for row in selected]),
        _metric(selected, "timeouts", [row["timeout"] for row in selected]),
        _metric(selected, "manual_review_required", [row["revisao_manual_necessaria"] for row in selected]),
    ]
    latency = _latency([float(row["total_time_ms"]) for row in selected if row["total_time_ms"] is not None])
    retrieval = _latency([float(row["retrieval_time_ms"]) for row in selected if row["retrieval_time_ms"] is not None])
    generation = _latency([float(row["generation_time_ms"]) for row in selected if row["generation_time_ms"] is not None])
    for prefix, values in (("latency_total_ms", latency), ("latency_retrieval_ms", retrieval), ("latency_generation_ms", generation)):
        for name, value in values.items():
            metrics.append({"metric": f"{prefix}_{name}", "value": value, "numerator": None, "denominator": len(selected) if value is not None else 0, "percentage": None, "measurable": value is not None, "notes": ""})
    for item in metrics:
        item["scope"] = scope
    return metrics


def _flat(row: dict[str, Any]) -> dict[str, Any]:
    excluded = {
        "raw_record",
        "judge_raw",
        "retrieval_results",
        "stage_one_retrieval_results",
        "context_results",
        "sources",
        "settings",
    }
    result: dict[str, Any] = {}
    for key, value in row.items():
        if key in excluded:
            continue
        result[key] = json.dumps(value, ensure_ascii=False) if isinstance(value, (list, dict)) else value
    result["documents_retrieved"] = json.dumps(row["retrieved_documents"], ensure_ascii=False)
    result["pages_retrieved"] = json.dumps(row["retrieved_pages"], ensure_ascii=False)
    result["retrieval_results_json"] = json.dumps(row["retrieval_results"], ensure_ascii=False)
    result["stage_one_retrieval_results_json"] = json.dumps(
        row["stage_one_retrieval_results"], ensure_ascii=False
    )
    result["context_results_json"] = json.dumps(
        row["context_results"], ensure_ascii=False
    )
    result["sources_json"] = json.dumps(row["sources"], ensure_ascii=False)
    return result


def _write_xlsx(path: Path, rows: list[dict[str, Any]], summary: list[dict[str, Any]]) -> None:
    workbook = Workbook()
    details = workbook.active
    details.title = "resultados_detalhados"
    flat = [_flat(row) for row in rows]
    headers = list(flat[0]) if flat else ["id"]
    details.append(headers)
    for row in flat:
        details.append([row.get(header) for header in headers])
    metrics = workbook.create_sheet("resumo_metricas")
    metric_headers = list(summary[0]) if summary else ["metric"]
    metrics.append(metric_headers)
    for row in summary:
        metrics.append([row.get(header) for header in metric_headers])
    for sheet in (details, metrics):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    for index in range(1, len(headers) + 1):
        details.column_dimensions[details.cell(1, index).column_letter].width = 24
    workbook.save(path)
    # Valida que o arquivo gerado pode ser reaberto e contém todas as linhas.
    check = load_workbook(path, read_only=True)
    assert check["resultados_detalhados"].max_row - 1 == len(rows)
    check.close()


def _display(item: dict[str, Any]) -> str:
    if item["value"] is None:
        return "não mensurável"
    if item["percentage"] is not None:
        if item["numerator"] is None:
            return f"{item['percentage']:.1f}% (N={item['denominator']})"
        return f"{item['percentage']:.1f}% ({item['numerator']}/{item['denominator']})"
    return f"{item['value']:.4f} (N={item['denominator']})"


def _latex_escape(value: str) -> str:
    for old, new in (("\\", r"\textbackslash{}"), ("_", r"\_"), ("%", r"\%"), ("&", r"\&"), ("#", r"\#")):
        value = value.replace(old, new)
    return value


def _write_documents(
    output_dir: Path,
    rows: list[dict[str, Any]],
    summary: list[dict[str, Any]],
    metadata: dict[str, Any],
) -> None:
    overall = {item["metric"]: item for item in summary if item["scope"] == "overall"}
    selected_names = [
        "hit_rate_at_1", f"hit_rate_at_{metadata['top_k']}", "mrr", "page_exact",
        "page_tolerance_1", "citation_document_correct", "citation_page_exact",
        "factual_correct", "mean_completeness", "mean_faithfulness", "mean_relevance",
        "fully_correct", "correct_refusal", "improper_refusal", "improper_answer",
        "decision_accuracy", "decision_precision", "decision_recall", "decision_f1",
        "always_answer_accuracy", "valid_inline_source", "correct_refusal_without_source",
        "latency_total_ms_mean", "latency_total_ms_p95", "latency_total_ms_p99",
        "errors", "timeouts", "manual_review_required",
    ]
    table = "\n".join(
        f"| {name} | {_display(overall[name])} | {overall[name]['notes']} |"
        for name in selected_names
    )
    report = f"""# Relatório de avaliação completa do RAG

> **Configuração avaliada:** {metadata['retrieval_provider']}. Este artefato registra
> uma execução específica e não deve ser atribuído a outro recuperador.

## Metodologia

- Benchmark congelado: `{metadata['benchmark_sha256']}`.
- Amostra: {metadata['question_count']} perguntas ({metadata['answerable_count']} com resposta no acervo e {metadata['unanswerable_count']} sem resposta no acervo).
- Execuções por pergunta: 1. Estabilidade entre repetições: não mensurável nesta execução principal.
- Recuperação reportada: k={metadata['top_k']}; pool interno de candidatos={metadata['candidate_pool_size']}.
- Página utilizada: página física do PDF (`page_start/page_end` do Chroma). A página impressa foi preservada separadamente quando disponível.
- Juiz de qualidade: `{metadata['judge_model']}`, temperatura 0, prompt versionado {JUDGE_VERSION}.
- O conjunto foi tratado como fixo durante esta execução; nenhum campo do gabarito foi modificado.

## Fórmulas

- Hit Rate@k = perguntas com o documento esperado entre as k primeiras / perguntas com resposta no acervo.
- RR = 1/posição do primeiro documento esperado; MRR = média dos RR.
- Completude = fatos obrigatórios corretamente mencionados / fatos obrigatórios.
- Resposta plenamente correta exige correção factual, completude 1, fidelidade 1, relevância 1, sustentação da citação 1 e ausência de recusa.

## Resultados gerais

| Métrica | Resultado | Observação |
|---|---:|---|
{table}

## Validação e limitações

- IDs duplicados: nenhum.
- Perguntas ignoradas: nenhuma.
- Erros e timeouts permanecem no denominador operacional e são apresentados explicitamente.
- Avaliação factual e evidência suficiente usam juiz LLM e não substituem validação humana.
- Foram marcadas {sum(row['revisao_manual_necessaria'] for row in rows)} respostas para revisão manual, incluindo todos os casos reprovados/ambíguos e uma amostra sistemática de 20% dos demais.
- A comparação com execuções anteriores não é válida neste relatório porque elas não utilizaram exatamente o mesmo texto das 100 perguntas.
- Uma execução por pergunta é o resultado principal; três repetições ficaram fora do escopo por custo e tempo, logo estabilidade não é mensurável aqui.
"""
    (output_dir / "relatorio_avaliacao.md").write_text(report, encoding="utf-8")
    latex_rows = "\n".join(
        f"{_latex_escape(name)} & {_latex_escape(_display(overall[name]))} \\\\" for name in selected_names
    )
    latex = """\\begin{table}[htbp]
\\centering
\\caption{{Resultados da execução com {metadata['retrieval_provider']}}}
\\label{tab:avaliacao-rag}
\\begin{tabular}{lr}
\\hline
Métrica & Resultado \\\\
\\hline
""" + latex_rows + "\n\\hline\n\\end{tabular}\n\\end{table}\n"
    (output_dir / "tabela_resultados_latex.tex").write_text(latex, encoding="utf-8")
    readme = f"""# Reprodução da avaliação

Configuração registrada neste diretório: **{metadata['retrieval_provider']}**.

Pré-requisitos: API ativa em `127.0.0.1:8000`, 43 PDFs processados, Chroma carregado e provedor LLM configurado.

Benchmark congelado: `{metadata['benchmark_sha256']}`.

```bash
python -m evaluation.validate_benchmark --benchmark evaluation/benchmark/benchmark_approved.json
python -m evaluation.freeze_config
python -m evaluation.run --benchmark evaluation/benchmark/benchmark_approved.json --split development --output evaluation/results/complete/development_run.json
python -m evaluation.run --benchmark evaluation/benchmark/benchmark_approved.json --split final --confirm-final --output evaluation/results/complete/final_run.json
python -m evaluation.complete --benchmark evaluation/benchmark/benchmark_approved.json --development-run evaluation/results/complete/development_run.json --final-run evaluation/results/complete/final_run.json --output-dir evaluation/results/complete
```

O avaliador grava checkpoint do juiz em `judge_checkpoint.jsonl`. Remova esse arquivo somente se desejar pagar e executar novamente todas as avaliações do juiz. Para três repetições, execute cada split três vezes com nomes distintos e agregue por ID; esta execução oficial utilizou uma repetição.
"""
    (output_dir / "README_avaliacao.md").write_text(readme, encoding="utf-8")


def evaluate_complete(
    benchmark_path: Path,
    development_run_path: Path,
    final_run_path: Path,
    output_dir: Path,
    use_llm_judge: bool = True,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    benchmark_payload = read_json(benchmark_path)
    questions = benchmark_payload["questions"]
    by_id = {question["id"]: question for question in questions}
    if len(by_id) != len(questions):
        raise ValueError("benchmark possui IDs duplicados")
    development = read_json(development_run_path)
    final = read_json(final_run_path)
    raw_records = list(development["records"]) + list(final["records"])
    ids = [record["id"] for record in raw_records]
    if len(raw_records) != 100 or len(set(ids)) != 100 or set(ids) != set(by_id):
        raise ValueError("execução deve conter exatamente os 100 IDs únicos do benchmark")
    _hydrate_sources(raw_records)
    output_dir.mkdir(parents=True, exist_ok=True)
    checkpoint_path = output_dir / "judge_checkpoint.jsonl"
    checkpoint: dict[str, dict[str, Any]] = {}
    if checkpoint_path.is_file():
        for line in checkpoint_path.read_text(encoding="utf-8").splitlines():
            if line.strip():
                item = json.loads(line)
                checkpoint[item["id"]] = item
    provider = create_llm_provider(get_settings()) if use_llm_judge else None
    rows: list[dict[str, Any]] = []
    for index, record in enumerate(raw_records, 1):
        question = by_id[record["id"]]
        judge = raw = error = None
        if question["answerable"]:
            if not use_llm_judge:
                error = "Não mensurável sem avaliação humana ou autorização do juiz LLM externo."
            else:
                cached = checkpoint.get(record["id"])
                if cached:
                    judge, raw, error = cached.get("judge"), cached.get("raw"), cached.get("error")
                else:
                    try:
                        judge, raw = _judge_one(provider, question, record)
                    except Exception as exc:  # checkpointa falha sem expor credenciais
                        error = f"{type(exc).__name__}: {exc}"
                    item = {"id": record["id"], "judge": judge, "raw": raw, "error": error}
                    with checkpoint_path.open("a", encoding="utf-8") as handle:
                        handle.write(json.dumps(item, ensure_ascii=False) + "\n")
        rows.append(_derived_record(question, record, judge, raw, error, 1))
        print(f"avaliado {index}/100: {record['id']}", flush=True)
    summary = [
        item
        for scope in ("overall", "development", "final")
        for item in calculate_summary(rows, scope)
    ]
    with (output_dir / "resultados_brutos.jsonl").open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")
    with (output_dir / "resumo_metricas.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(summary[0]))
        writer.writeheader()
        writer.writerows(summary)
    _write_xlsx(output_dir / "resultados_detalhados.xlsx", rows, summary)
    settings = rows[0]["settings"]
    metadata = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "benchmark_sha256": _sha256(benchmark_path),
        "question_count": len(rows),
        "answerable_count": sum(row["answerable"] for row in rows),
        "unanswerable_count": sum(not row["answerable"] for row in rows),
        "top_k": settings.get("top_k"),
        "candidate_pool_size": settings.get("candidate_pool_size"),
        "retrieval_provider": settings.get(
            "retrieval_provider",
            "recuperação densa (execução histórica)",
        ),
        "judge_model": get_settings().llm_model if use_llm_judge else "não executado",
        "judge_prompt_sha256": hashlib.sha256((JUDGE_SYSTEM_PROMPT + JUDGE_USER_TEMPLATE).encode()).hexdigest(),
    }
    (output_dir / "experiment_manifest.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    _write_documents(output_dir, rows, summary, metadata)
    return rows, summary


def main() -> None:
    parser = argparse.ArgumentParser(description="Gera a avaliação completa das 100 perguntas")
    parser.add_argument("--benchmark", type=Path, required=True)
    parser.add_argument("--development-run", type=Path, required=True)
    parser.add_argument("--final-run", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument(
        "--skip-llm-judge",
        action="store_true",
        help="Gera apenas métricas determinísticas e marca qualidade como não mensurável.",
    )
    args = parser.parse_args()
    rows, _ = evaluate_complete(
        args.benchmark,
        args.development_run,
        args.final_run,
        args.output_dir,
        use_llm_judge=not args.skip_llm_judge,
    )
    print(f"Avaliação completa gerada em {args.output_dir}: {len(rows)} registros.")


if __name__ == "__main__":
    main()
