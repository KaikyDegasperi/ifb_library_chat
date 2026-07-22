"""Relatórios autocontidos em JSON, CSV, XLSX e HTML."""

import argparse
import json
from html import escape
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from evaluation.io import read_json, write_csv, write_json
from evaluation.metrics import calculate_metrics, grouped_metrics
from evaluation.models import RunOutput


def build_report(run: RunOutput) -> dict[str, Any]:
    records = [record.model_dump(mode="json") for record in run.records]
    return {
        "report_version": "1.0",
        "split": run.split,
        "benchmark_version": run.benchmark_version,
        "run_period": {"started_at": run.started_at, "completed_at": run.completed_at},
        "settings": run.settings,
        "metrics": calculate_metrics(records),
        "metrics_by_question_type": grouped_metrics(records, "question_type"),
        "metrics_by_difficulty": grouped_metrics(records, "difficulty"),
        "metrics_by_answerability": grouped_metrics(records, "answerable"),
        "records": records,
    }


def _flat_record(record: dict[str, Any]) -> dict[str, Any]:
    return {
        **record,
        "required_facts": json.dumps(record.get("required_facts", []), ensure_ascii=False),
        "expected_pages": json.dumps(record.get("expected_pages", []), ensure_ascii=False),
        "retrieved_documents": json.dumps(record.get("retrieved_documents", []), ensure_ascii=False),
        "retrieved_pages": json.dumps(record.get("retrieved_pages", []), ensure_ascii=False),
        "similarities": json.dumps(record.get("similarities", []), ensure_ascii=False),
        "returned_passages": json.dumps(record.get("returned_passages", []), ensure_ascii=False),
        "sources": json.dumps(record.get("sources", []), ensure_ascii=False),
        "retrieval_results": json.dumps(record.get("retrieval_results", []), ensure_ascii=False),
        "settings": json.dumps(record.get("settings", {}), ensure_ascii=False),
    }


def write_xlsx(path: Path, report: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "summary"
    summary.append(["Métrica", "Valor"])
    for key, value in report["metrics"].items():
        summary.append([key, value])
    summary.append([])
    summary.append(["Configuração", "Valor"])
    for key, value in report["settings"].items():
        summary.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (list, dict)) else value])
    details = workbook.create_sheet("details")
    rows = [_flat_record(record) for record in report["records"]]
    columns = list(rows[0]) if rows else ["id"]
    details.append(columns)
    for row in rows:
        details.append([row.get(column) for column in columns])
    for sheet in (summary, details):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    score_fields = {"correctness", "faithfulness", "completeness", "citation_quality", "human_notes"}
    for index, column in enumerate(columns, 1):
        details.column_dimensions[details.cell(1, index).column_letter].width = 26
        if column in score_fields:
            for cells in details.iter_cols(min_col=index, max_col=index, min_row=2):
                cells[0].fill = PatternFill("solid", fgColor="FFF2CC")
    for sheet_name, group_key in (("by_type", "metrics_by_question_type"), ("by_difficulty", "metrics_by_difficulty"), ("by_answerability", "metrics_by_answerability")):
        grouped = workbook.create_sheet(sheet_name)
        metric_names = list(next(iter(report[group_key].values()), {}).keys())
        grouped.append(["group", *metric_names])
        for name, values in report[group_key].items():
            grouped.append([name, *(values.get(metric) for metric in metric_names)])
        for cell in grouped[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
    workbook.save(path)


def _metric_bars(metrics: dict[str, Any]) -> str:
    keys = ["document_recall_at_1", "document_recall_at_3", "document_recall_at_k", "page_recall_at_k"]
    return "".join(
        f"<div class='bar-row'><span>{escape(key)}</span><div class='bar'><i style='width:{float(metrics[key])*100:.1f}%'></i></div><b>{float(metrics[key]):.1%}</b></div>"
        for key in keys
    )


def render_html(report: dict[str, Any]) -> str:
    metrics = report["metrics"]
    type_rows = "".join(
        f"<tr><td>{escape(name)}</td><td>{values['question_count']}</td><td>{values['document_recall_at_k']:.1%}</td><td>{values['correct_refusal_rate']:.1%}</td></tr>"
        for name, values in report["metrics_by_question_type"].items()
    )
    detail_rows = []
    for record in report["records"]:
        detail_rows.append(
            "<tr><td>{id}</td><td>{question}</td><td>{expected}</td><td>{produced}</td><td>{expected_source}</td><td>{retrieved}</td><td>{time}</td><td>{error}</td><td>{scores}</td></tr>".format(
                id=escape(record["id"]), question=escape(record["question"]),
                expected=escape(record["expected_answer"]), produced=escape(record["produced_answer"]),
                expected_source=escape(f"{record.get('expected_document') or '—'} / {record.get('expected_pages', [])}"),
                retrieved=escape(str(record.get("retrieved_documents", []))), time=record.get("total_time_ms", 0),
                error=escape(str(record.get("error") or "")),
                scores=escape(str({key: record.get(key) for key in ("correctness", "faithfulness", "completeness", "citation_quality")})),
            )
        )
    refusal_correct = float(metrics["correct_refusal_rate"]) * 100
    refusal_wrong = float(metrics["improper_refusal_rate"]) * 100
    times = [float(record.get("total_time_ms", 0)) for record in report["records"]]
    maximum_time = max(times, default=1.0) or 1.0
    time_bars = "".join(
        f"<i title='{value:.0f} ms' style='height:{max(3, value / maximum_time * 100):.1f}%'></i>"
        for value in times
    ) or "<span>Sem execuções</span>"
    score_bars = []
    for field in ("correctness", "faithfulness"):
        values = [record[field] for record in report["records"] if record.get(field) is not None]
        average = sum(values) / len(values) if values else 0.0
        score_bars.append(f"<div class='bar-row'><span>{field}</span><div class='bar'><i style='width:{average/2*100:.1f}%'></i></div><b>{average:.2f}/2</b></div>")
    settings_rows = "".join(f"<tr><td>{escape(str(key))}</td><td>{escape(str(value))}</td></tr>" for key, value in report["settings"].items())
    return f"""<!doctype html><html lang='pt-BR'><meta charset='utf-8'><title>Relatório de avaliação RAG</title>
<style>body{{font:14px system-ui;margin:2rem;color:#17202a}}.cards{{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:1rem}}.card,section{{border:1px solid #d5dbe3;border-radius:10px;padding:1rem;margin:1rem 0}}.card strong{{font-size:1.5rem;display:block}}.bar-row{{display:grid;grid-template-columns:200px 1fr 70px;gap:.7rem;align-items:center;margin:.6rem 0}}.bar{{height:18px;background:#e9eef4;border-radius:9px;overflow:hidden}}.bar i{{height:100%;display:block;background:#16734b}}.hist{{height:130px;display:flex;align-items:end;gap:3px;border-bottom:1px solid #9aa;padding-top:1rem}}.hist i{{background:#3978a8;min-width:8px;flex:1}}table{{border-collapse:collapse;width:100%}}th,td{{border:1px solid #ddd;padding:.45rem;text-align:left;vertical-align:top}}th{{background:#163a5f;color:#fff}}.scroll{{overflow:auto;max-height:70vh}}small{{color:#536273}}</style>
<h1>Relatório de avaliação do RAG</h1><small>Split {escape(report['split'])} · benchmark {escape(report['benchmark_version'])}</small>
<div class='cards'><div class='card'>Perguntas<strong>{metrics['question_count']}</strong></div><div class='card'>MRR<strong>{metrics['mrr']:.1%}</strong></div><div class='card'>Tempo médio<strong>{metrics['mean_time_ms']:.0f} ms</strong></div><div class='card'>Erros / timeouts<strong>{metrics['error_count']} / {metrics['timeout_count']}</strong></div></div>
<section><h2>Recuperação</h2>{_metric_bars(metrics)}</section>
<section><h2>Recuperação por tipo</h2><table><tr><th>Tipo</th><th>N</th><th>Recall@k</th><th>Recusa correta</th></tr>{type_rows}</table></section>
<section><h2>Recusas</h2><div class='bar-row'><span>Corretas</span><div class='bar'><i style='width:{refusal_correct:.1f}%'></i></div><b>{refusal_correct:.1f}%</b></div><div class='bar-row'><span>Indevidas</span><div class='bar'><i style='width:{refusal_wrong:.1f}%;background:#b33'></i></div><b>{refusal_wrong:.1f}%</b></div></section>
<section><h2>Distribuição dos tempos</h2><div class='hist'>{time_bars}</div><p>Média {metrics['mean_time_ms']:.0f} ms · mediana {metrics['median_time_ms']:.0f} ms · p95 {metrics['p95_time_ms']:.0f} ms</p></section>
<section><h2>Notas humanas</h2>{''.join(score_bars)}<p>Os campos de correção, fidelidade, completude, qualidade de citação e observações ficam editáveis (amarelos) no XLSX. Valores ausentes não são substituídos por avaliação automática.</p></section>
<section><h2>Configuração</h2><table><tr><th>Parâmetro</th><th>Valor</th></tr>{settings_rows}</table></section>
<section><h2>Execuções detalhadas</h2><div class='scroll'><table><tr><th>ID</th><th>Pergunta</th><th>Esperada</th><th>Obtida</th><th>Fonte esperada</th><th>Documentos recuperados</th><th>ms</th><th>Erro</th><th>Notas</th></tr>{''.join(detail_rows)}</table></div></section></html>"""


def generate_reports(run_path: Path, output_dir: Path) -> dict[str, Any]:
    run = RunOutput.model_validate(read_json(run_path))
    report = build_report(run)
    output_dir.mkdir(parents=True, exist_ok=True)
    write_json(output_dir / "report.json", report)
    write_csv(output_dir / "report.csv", [_flat_record(record) for record in report["records"]])
    write_xlsx(output_dir / "report.xlsx", report)
    (output_dir / "report.html").write_text(render_html(report), encoding="utf-8")
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description="Gera relatórios de uma execução")
    parser.add_argument("--run", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, default=Path("evaluation/results"))
    args = parser.parse_args()
    report = generate_reports(args.run, args.output_dir)
    print(f"Relatórios gerados em {args.output_dir}; {report['metrics']['question_count']} registros.")


if __name__ == "__main__":
    main()
