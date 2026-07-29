"""Leitura, escrita e representação tabular dos artefatos."""

import csv
import json
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from evaluation.models import Benchmark, BenchmarkQuestion


BENCHMARK_COLUMNS = [
    "id", "question", "answerable", "question_type", "difficulty",
    "expected_answer", "required_facts", "expected_document", "expected_pages",
    "printed_pages", "evidence", "section", "split", "status", "rationale",
    "unanswerable_reason", "review_notes", "document_link",
]


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def read_benchmark(path: Path) -> Benchmark:
    return Benchmark.model_validate(read_json(path))


def write_benchmark(path: Path, benchmark: Benchmark) -> None:
    write_json(path, benchmark.model_dump(mode="json"))


def _join(values: list[Any]) -> str:
    return " | ".join(str(value) for value in values)


def question_row(question: BenchmarkQuestion, documents_dir: Path) -> dict[str, Any]:
    data = question.model_dump(mode="json")
    data["required_facts"] = _join(question.required_facts)
    data["expected_pages"] = _join(question.expected_pages)
    data["printed_pages"] = _join(question.printed_pages)
    data["document_link"] = (
        str((documents_dir / question.expected_document).resolve())
        if question.expected_document else ""
    )
    return data


def write_benchmark_xlsx(path: Path, benchmark: Benchmark, documents_dir: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    metadata = workbook.active
    metadata.title = "metadata"
    metadata.append(["benchmark_version", benchmark.benchmark_version])
    metadata.append(["seed", benchmark.seed])
    metadata.append(
        ["excluded_documents", json.dumps(benchmark.excluded_documents, ensure_ascii=False)]
    )
    metadata.sheet_state = "hidden"
    sheet = workbook.create_sheet()
    sheet.title = "questions"
    sheet.append(BENCHMARK_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for question in benchmark.questions:
        row = question_row(question, documents_dir)
        sheet.append([row.get(column, "") for column in BENCHMARK_COLUMNS])
        link_cell = sheet.cell(sheet.max_row, BENCHMARK_COLUMNS.index("document_link") + 1)
        if link_cell.value:
            link_cell.hyperlink = f"file://{link_cell.value}"
            link_cell.style = "Hyperlink"
    editable = {"question", "expected_answer", "required_facts", "status", "review_notes"}
    for index, column in enumerate(BENCHMARK_COLUMNS, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = (
            45 if column in {"question", "expected_answer", "evidence", "review_notes"} else 20
        )
        if column in editable:
            for cell in sheet.iter_cols(min_col=index, max_col=index, min_row=2):
                cell[0].fill = PatternFill("solid", fgColor="FFF2CC")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    guide = workbook.create_sheet("instructions")
    guide.append(["Revisão humana obrigatória"])
    guide.append(["Confira o PDF e as páginas reais (base 1) antes de aprovar."])
    guide.append(["Edite os campos amarelos e use approved, rejected ou pending_review."])
    guide.append(
        ["Perguntas com resposta no acervo exigem fonte verificável para aprovação."]
    )
    workbook.save(path)


def benchmark_from_xlsx(path: Path, seed: int = 42) -> Benchmark:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook["questions"]
    benchmark_version = "1.0"
    excluded_documents: dict[str, str] = {}
    if "metadata" in workbook.sheetnames:
        metadata = {
            str(row[0]): row[1]
            for row in workbook["metadata"].iter_rows(values_only=True)
            if row and row[0]
        }
        benchmark_version = str(metadata.get("benchmark_version") or "1.0")
        seed = int(metadata.get("seed") or seed)
        raw_exclusions = metadata.get("excluded_documents")
        if raw_exclusions:
            excluded_documents = json.loads(str(raw_exclusions))
    headers = [str(cell.value) for cell in sheet[1]]
    questions: list[BenchmarkQuestion] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values, strict=False))
        if not row.get("id"):
            continue
        row.pop("document_link", None)
        row["answerable"] = str(row.get("answerable", "")).casefold() in {"true", "1", "sim"}
        row["required_facts"] = _split(row.get("required_facts"), str)
        row["expected_pages"] = _split(row.get("expected_pages"), int)
        row["printed_pages"] = _split(row.get("printed_pages"), str)
        row["expected_document"] = row.get("expected_document") or None
        row["section"] = row.get("section") or None
        row["unanswerable_reason"] = row.get("unanswerable_reason") or None
        for field in ("question", "question_type", "difficulty", "expected_answer", "evidence", "split", "status", "rationale", "review_notes"):
            row[field] = "" if row.get(field) is None else str(row[field])
        questions.append(BenchmarkQuestion.model_validate(row))
    return Benchmark(
        benchmark_version=benchmark_version,
        seed=seed,
        excluded_documents=excluded_documents,
        questions=questions,
    )


def _split(value: Any, caster: type) -> list:
    if value is None or value == "":
        return []
    return [caster(part.strip()) for part in str(value).split("|") if part.strip()]


def write_csv(path: Path, rows: Iterable[dict[str, Any]], columns: list[str] | None = None) -> None:
    materialized = list(rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = columns or sorted({key for row in materialized for key in row})
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in materialized:
            writer.writerow({key: _json_cell(value) for key, value in row.items()})


def _json_cell(value: Any) -> Any:
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return value
